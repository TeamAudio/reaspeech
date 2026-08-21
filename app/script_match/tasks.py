import logging
import os
from typing import Dict, List
from app.celery_app import celery

logger = logging.getLogger(__name__)


def read_row_styles(spreadsheet_path: str) -> Dict[str, Dict[str, List[int]]]:
    """Return per-sheet row style info: 1-based "hidden" and "bold" row
    numbers (a row is bold when any non-empty cell in it has a bold font).

    pandas does not expose row visibility or styling, so this uses openpyxl
    directly. The workbook is passed as an open file object because
    load_workbook rejects paths without a recognized extension, and uploads
    arrive as extensionless temp files. Returns empty info for formats
    openpyxl cannot read (e.g. .xls).
    """
    styles: Dict[str, Dict[str, List[int]]] = {}
    try:
        import openpyxl
        with open(spreadsheet_path, 'rb') as f:
            workbook = openpyxl.load_workbook(f, read_only=False)
        for worksheet in workbook.worksheets:
            hidden = sorted(
                row_number
                for row_number, dimension in worksheet.row_dimensions.items()
                if dimension.hidden
            )
            bold = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None and cell.font and cell.font.bold:
                        bold.append(cell.row)
                        break
            styles[worksheet.title] = {"hidden": hidden, "bold": bold}
    except Exception as e:
        logger.warning(f"Could not read row style info: {e}")
    return styles

@celery.task(name="script_match.parse_spreadsheet", bind=True)
def parse_spreadsheet(
    self,
    spreadsheet_path: str,
    filename: str
):
    """
    Parse an uploaded spreadsheet into per-sheet cell data plus row style
    info (hidden and bold rows).

    Args:
        spreadsheet_path: Path to uploaded spreadsheet file
        filename: Original filename
    
    Returns:
        Parsed spreadsheet data structure
    """
    logger.info(f"Parsing spreadsheet: {filename}")
    
    try:
        from pandas import read_excel, isna

        self.update_state(
            state='PARSING',
            meta={'progress': {'units': 'files', 'total': 1, 'current': 0}}
        )
        
        sheet_dict = read_excel(
            spreadsheet_path, 
            verbose=True, 
            header=None, 
            sheet_name=None, 
            na_values=['', 'NaT']
        )
        
        row_styles = read_row_styles(spreadsheet_path)

        def make_sheet(name, sheet):
            sheet_styles = row_styles.get(name, {})
            return {
                "config": {
                    "name": name,
                    "column_count": sheet.shape[1],
                    "row_count": sheet.shape[0],
                    "hidden_rows": sheet_styles.get("hidden", []),
                    "bold_rows": sheet_styles.get("bold", []),
                },
                "data": [[str(v_) if not isna(v_) else '' for v_ in list(v)] for _, v in sheet.iterrows()]
            }
        
        sheets = [make_sheet(name, s) for name, s in sheet_dict.items()]
        
        result = {
            "config": {
                "file_type": "excel",
                "file_name": filename,
                "sheets": len(sheets),
            },
            "sheets": sheets
        }
        
        logger.info(f"Spreadsheet parsed successfully: {len(sheets)} sheets")
        return result
        
    except Exception as e:
        logger.error(f"Spreadsheet parsing error: {str(e)}")
        raise
    finally:
        # Clean up temporary file
        try:
            os.unlink(spreadsheet_path)
        except:
            pass
